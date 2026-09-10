"""
Extracción de texto de documentos DOCX (Word) con texto digital.

No incluye (todavía) OCR sobre imágenes incrustadas dentro del docx -- si
hiciera falta, se añadiría siguiendo el mismo patrón que en pdf_extractor.py.
Por ahora no era un requisito planteado para el caso de uso.
"""

import docx


def extraer_texto_docx(ruta_archivo: str) -> str:
    """
    Devuelve el texto del documento como una concatenación de:
      - los párrafos, en orden de aparición
      - el contenido de las tablas, fila por fila (celdas separadas por " | ")

    No distingue estilos (encabezados, negrita, etc.). Si el Agente técnico
    necesitara en el futuro identificar secciones por su estilo, se puede
    extender aquí sin afectar al resto del sistema.
    """
    documento = docx.Document(ruta_archivo)
    partes = []

    for parrafo in documento.paragraphs:
        if parrafo.text.strip():
            partes.append(parrafo.text.strip())

    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            if any(celdas):
                partes.append(" | ".join(celdas))

    return "\n".join(partes)