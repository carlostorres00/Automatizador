"""
Extracción de contenido tabular de archivos XLSX.

Devuelve los valores calculados (no las fórmulas) de cada hoja, como listas
de filas. Esto es suficiente para el Agente económico más adelante -- no
intentamos aquí interpretar qué fila es qué partida, eso es trabajo de
análisis, no de ingesta.
"""

import openpyxl


def extraer_tablas_xlsx(ruta_archivo: str) -> dict[str, list[list]]:
    """
    Devuelve un diccionario {nombre_hoja: filas}, donde cada fila es una lista
    de valores de celda (ya evaluados, no fórmulas -- usamos data_only=True).
    """
    libro = openpyxl.load_workbook(ruta_archivo, data_only=True)
    resultado: dict[str, list[list]] = {}
    for nombre_hoja in libro.sheetnames:
        hoja = libro[nombre_hoja]
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            # Se descartan filas completamente vacías, para no ensuciar el resultado.
            if any(celda is not None for celda in fila):
                filas.append(list(fila))
        resultado[nombre_hoja] = filas
    return resultado
